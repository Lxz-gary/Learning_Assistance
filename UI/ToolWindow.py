"""工具页：Excel 批量导入、全量导出、数据库备份。"""

import os
import shutil
import time
import tkinter as tk
from tkinter import filedialog

import customtkinter as ctk
from openpyxl import load_workbook

from Controller.DataBaseController import DataBaseController
from Utils.ExcelOperater import export_to_excel
from .theme import COLORS, font
from .widgets import build_card, make_btn, build_table_card, ask_dialog

db_controller = DataBaseController()

DB_FILE = "my_database.db"

# 每种类型：Excel 表头 -> 数据库字段顺序
IMPORT_SPECS = {
    "学生": {
        "headers": ["学生ID", "姓名", "班级", "年级", "电话号码", "邮箱"],
        "columns": ("id", "name", "class", "grade", "telephonenum", "email"),
        "labels": ["编号", "姓名", "班级", "年级", "电话", "邮箱"],
        "extract": lambda r: (str(r[0]), str(r[1]), str(r[2]), str(r[3]), str(r[4]), str(r[5])),
    },
    "考试": {
        "headers": ["考试ID", "考试名称", "考试日期"],
        "columns": ("exam_id", "name", "date"),
        "labels": ["考试编号", "考试名称", "考试日期"],
        "extract": lambda r: (str(r[0]), str(r[1]), str(r[2])),
    },
    "成绩": {
        "headers": ["学生ID", "考试ID", "科目", "成绩", "排名", "总人数", "赋分"],
        "columns": ("student_id", "exam_id", "subject", "initialscore", "rank", "headcount", "score"),
        "labels": ["学号", "考试编号", "科目", "初始分", "排名", "总人数", "赋分"],
        "extract": lambda r: (str(r[0]), str(r[1]), str(r[2]),
                              int(r[3]), int(r[4]), int(r[5]), int(r[6])),
    },
}


def _table_config(spec):
    return {
        "columns": spec["columns"],
        "headers": {c: (l, 110) for c, l in zip(spec["columns"], spec["labels"])},
        "card_bg": "sidebar_bg", "row_bg": "sidebar_bg", "row_alt": "content_bg",
        "heading_bg": "sidebar_active", "heading_fg": "Level1fg",
        "heading_act_bg": "accent", "heading_act_fg": "card_bg",
        "selected_bg": "accent", "selected_fg": "selected_fg",
        "h_scroll": True,
    }


class ToolWindow:
    def __init__(self, root: tk.Tk, content_frame: tk.Frame):
        self.root = root
        self.content_frame = content_frame
        self._path = None
        self._rows = []
        self._kind = "学生"

    # ------------------------------------------------------------------
    def _reset(self, title, subtitle):
        for widget in self.content_frame.winfo_children():
            widget.destroy()
        head = tk.Frame(self.content_frame, bg=COLORS["content_bg"])
        head.pack(fill="x", padx=20, pady=(18, 2))
        tk.Label(head, text=title, bg=COLORS["content_bg"], fg=COLORS["Level1fg"],
                 font=font(20, "bold"), anchor="w").pack(anchor="w")
        tk.Label(head, text=subtitle, bg=COLORS["content_bg"], fg=COLORS["Level2fg"],
                 font=font(11), anchor="w").pack(anchor="w", pady=(2, 0))

    # ==================================================================
    # 导入
    # ==================================================================
    def show_import(self):
        self._reset("导入 Excel", "选择与模版表头一致的文件，先预览再导入")

        bar_card = build_card(self.content_frame)
        bar_card.pack(fill="x", padx=20, pady=(14, 0))
        bar = tk.Frame(bar_card, bg=COLORS["content_bg"])
        bar.pack(fill="x", padx=14, pady=12)

        tk.Label(bar, text="数据类型", bg=COLORS["content_bg"], fg=COLORS["Level2fg"],
                 font=font(12)).pack(side="left", padx=(0, 8))
        self._type_box = ctk.CTkComboBox(
            bar, values=list(IMPORT_SPECS.keys()), width=120, height=28,
            border_width=0, state="readonly", command=self._on_type_change,
            fg_color=COLORS["entry_bg"], text_color=COLORS["Level1fg"],
            button_color=COLORS["btn_bg"], button_hover_color=COLORS["btn_hover"],
            dropdown_fg_color=COLORS["sidebar_bg"],
            dropdown_text_color=COLORS["Level1fg"],
            dropdown_hover_color=COLORS["btn_hover"],
        )
        self._type_box.pack(side="left")
        self._type_box.set(self._kind)

        self._path_label = tk.Label(
            bar, text="尚未选择文件", bg=COLORS["content_bg"], fg=COLORS["Level2fg"],
            font=font(11), anchor="w",
        )
        self._path_label.pack(side="left", padx=14, fill="x", expand=True)

        make_btn(bar, "选择文件", self._pick_file, kind="secondary",
                 size=11, padx=16, pady=7).pack(side="right")
        make_btn(bar, "开始导入", self._do_import, kind="primary",
                 size=12, padx=20, pady=8).pack(side="right", padx=(0, 8))

        self._hint = tk.Label(
            self.content_frame, text="表头需与模版一致："
            + " / ".join(IMPORT_SPECS[self._kind]["headers"]),
            bg=COLORS["content_bg"], fg=COLORS["Level2fg"], font=font(10), anchor="w",
        )
        self._hint.pack(fill="x", padx=20, pady=(10, 6))

        # 预览表
        self._preview_holder = tk.Frame(self.content_frame, bg=COLORS["content_bg"])
        self._preview_holder.pack(fill="both", expand=True, padx=20, pady=(0, 16))
        self._render_preview()

    def _on_type_change(self, value):
        self._kind = value
        self._rows = []
        self._path = None
        if hasattr(self, "_path_label"):
            self._path_label.configure(text="尚未选择文件")
        if hasattr(self, "_hint"):
            self._hint.configure(
                text="表头需与模版一致：" + " / ".join(IMPORT_SPECS[self._kind]["headers"])
            )
        self._render_preview()

    def _render_preview(self):
        for w in self._preview_holder.winfo_children():
            w.destroy()
        spec = IMPORT_SPECS[self._kind]
        card, tree = build_table_card(self._preview_holder, _table_config(spec), height=10)
        card.pack(fill="both", expand=True)
        for i, row in enumerate(self._rows[:50]):
            tree.insert("", "end", values=row,
                        tags=("evenrow" if i % 2 == 0 else "oddrow",))
        if not self._rows:
            tree.insert("", "end", values=("—",) * len(spec["columns"]),
                        tags=("empty",))

    def _pick_file(self):
        path = filedialog.askopenfilename(
            title="选择 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
        )
        if not path:
            return
        try:
            rows = self._read_excel(path, self._kind)
        except Exception as exc:
            ask_dialog(self.root, "读取失败", str(exc), ok_text="知道了", cancel_text=None)
            return
        self._path = path
        self._rows = rows
        self._path_label.configure(text=os.path.basename(path))
        self._hint.configure(
            text=f"已读取 {len(rows):,} 行，下面预览前 50 行；确认无误后点「开始导入」"
        )
        self._render_preview()

    def _read_excel(self, path, kind):
        spec = IMPORT_SPECS[kind]
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            raise ValueError("文件是空的")

        header = [str(c).strip() if c is not None else "" for c in rows[0][:len(spec["headers"])]]
        if header != spec["headers"]:
            raise ValueError(
                f"表头不匹配。\n期望：{' / '.join(spec['headers'])}\n实际：{' / '.join(header)}"
            )

        out = []
        for line in rows[1:]:
            if line is None or all(c is None or str(c).strip() == "" for c in line):
                continue
            try:
                out.append(spec["extract"](line))
            except Exception:
                continue        # 单行格式错误就跳过，最后统计里体现不出来，但不会阻断导入
        return out

    def _do_import(self):
        if not self._rows:
            ask_dialog(self.root, "提示", "请先选择一个 Excel 文件。",
                       ok_text="知道了", cancel_text=None)
            return
        self._hint.configure(text="正在导入，请稍候…", fg=COLORS["Level2fg"])
        self.root.update_idletasks()

        started = time.perf_counter()
        inserted, duplicated, invalid = db_controller.import_rows(self._kind, self._rows)
        cost = time.perf_counter() - started

        text = (f"导入完成：成功 {inserted:,} 条，重复跳过 {duplicated:,} 条，"
                f"无效跳过 {invalid:,} 条，耗时 {cost:.1f}s")
        self._hint.configure(text=text, fg=COLORS["ok_fg"] if inserted else COLORS["warn_fg"])
        ask_dialog(self.root, "导入完成", text, ok_text="好的", cancel_text=None)

    # ==================================================================
    # 导出
    # ==================================================================
    def show_export(self):
        self._reset("导出数据", "把某一张表的全部数据导出成带样式的 Excel")

        card = build_card(self.content_frame)
        card.pack(fill="x", padx=20, pady=(14, 0))
        bar = tk.Frame(card, bg=COLORS["content_bg"])
        bar.pack(fill="x", padx=14, pady=12)

        tk.Label(bar, text="导出内容", bg=COLORS["content_bg"], fg=COLORS["Level2fg"],
                 font=font(12)).pack(side="left", padx=(0, 8))
        self._export_box = ctk.CTkComboBox(
            bar, values=["学生", "考试", "成绩"], width=120, height=28,
            border_width=0, state="readonly",
            fg_color=COLORS["entry_bg"], text_color=COLORS["Level1fg"],
            button_color=COLORS["btn_bg"], button_hover_color=COLORS["btn_hover"],
            dropdown_fg_color=COLORS["sidebar_bg"],
            dropdown_text_color=COLORS["Level1fg"],
            dropdown_hover_color=COLORS["btn_hover"],
        )
        self._export_box.pack(side="left")
        self._export_box.set("学生")

        make_btn(bar, "选择保存位置并导出", self._do_export, kind="primary",
                 size=12, padx=20, pady=8).pack(side="right")

        self._export_hint = tk.Label(
            self.content_frame, text="", bg=COLORS["content_bg"],
            fg=COLORS["Level2fg"], font=font(10), anchor="w",
        )
        self._export_hint.pack(fill="x", padx=20, pady=(10, 0))

    def _do_export(self):
        kind = self._export_box.get()
        path = filedialog.asksaveasfilename(
            title="保存文件", defaultextension=".xlsx",
            initialfile=f"{kind}全部数据.xlsx",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
        )
        if not path:
            return

        if kind == "学生":
            headers = ["编号", "姓名", "班级", "年级", "电话", "邮箱", "创建时间", "修改时间"]
            rows = db_controller.db_manager.query_students(None, None, None, None, None, None, None, None)
        elif kind == "考试":
            headers = ["考试编号", "考试名称", "考试日期", "创建时间", "修改时间"]
            rows = db_controller.db_manager.query_exams(None, None, None, None, None)
        else:
            headers = ["学号", "考试编号", "科目", "初始分", "成绩", "排名", "总人数", "创建时间", "修改时间"]
            rows = db_controller.db_manager.query_scores(None, None, None, None, None, None, None, None, None)

        self._export_hint.configure(text=f"正在导出 {len(rows):,} 条…")
        self.root.update_idletasks()
        export_to_excel(headers=headers, rows=rows, filepath=path,
                        sheet_name="数据", title=f"{kind}全部数据")
        self._export_hint.configure(
            text=f"✓ 已导出 {len(rows):,} 条到 {os.path.basename(path)}", fg=COLORS["ok_fg"]
        )

    # ==================================================================
    # 备份（动作，不占页面）
    # ==================================================================
    def backup_database(self):
        if not os.path.exists(DB_FILE):
            ask_dialog(self.root, "备份失败", f"找不到数据库文件：{DB_FILE}",
                       ok_text="知道了", cancel_text=None)
            return
        backup_dir = "backups"
        os.makedirs(backup_dir, exist_ok=True)
        name = f"my_database_{time.strftime('%Y%m%d_%H%M%S')}.db"
        dest = os.path.join(backup_dir, name)
        shutil.copyfile(DB_FILE, dest)
        size = os.path.getsize(dest) / 1024 / 1024
        ask_dialog(self.root, "备份完成", f"已备份到：\n{dest}\n（{size:.1f} MB）",
                   ok_text="好的", cancel_text=None)
